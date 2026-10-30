"""把 GTDS（HUST_GTDS 运行模拟）的 Excel 数据转换为算法输入目录。

用到三类文件：
- 日历负荷 L2030-8760-Ax.xlsx：365 天 × 24 h 标幺值；
- 电站出力一览表：G 表为发电出力、U 表为调峰弃电，可发出力 = G + U；
- 机组参数表：系统表、负荷表、节点表、电站表、机组表。

GTDS 电站类型：300 煤电，320/325 燃气，360 核电，370/373 水电，380 抽蓄，
385 新型储能，390/392 风电，391 光伏，393 其他。
"""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

COAL, GAS, NUCLEAR, OTHER = {300}, {320, 325}, {360}, {393}
HYDRO, PUMPED, BATTERY = {370, 373}, {380}, {385}
WIND, SOLAR = {390, 392}, {391}
GAS_PRIORITY = 400.0     # 燃气机组没有煤耗数据，排在所有煤电之后
OTHER_PRIORITY = 0.0     # "其他"电源按必开处理


def _sheet_rows(ws):
    return [r for r in ws.iter_rows(values_only=True)]


def _table(wb, name) -> list[dict]:
    rows = _sheet_rows(wb[name])
    head = rows[0]
    out = []
    for r in rows[1:]:
        if r[0] is None or r[1] is None:
            continue
        d = {}
        for h, v in zip(head, r):            # 表头有重名列（如负荷表两列"年Pmax"），取第一列
            if h is not None and h not in d:
                d[h] = v
        out.append(d)
    return out


def read_daily_curve(path) -> np.ndarray:
    """日历负荷：首列日期，后 24 列 L1..L24。"""
    rows = _sheet_rows(openpyxl.load_workbook(path, data_only=True).worksheets[0])
    return np.array([[float(v) for v in r[1:25]] for r in rows[1:] if r[0]], float).ravel()


def read_plant_sheets(path) -> dict[int, np.ndarray]:
    """电站出力一览表 → {电站ID: 可发出力（G+U）}。按表头标题里的 "G<ID>_" 识别电站。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    out: dict[int, np.ndarray] = {}
    for ws in wb.worksheets:
        rows = _sheet_rows(ws)
        m = re.search(r"G(\d+)_", str(rows[0][1]))
        if not m:
            continue
        hdr = rows[3]
        cols = [i for i, c in enumerate(hdr) if isinstance(c, str) and re.fullmatch(r"\d+时", c.strip())]
        data = [r for r in rows if isinstance(r[0], str) and re.match(r"\d{4}/", r[0].strip())]
        s = np.array([[float(r[i] or 0) for i in cols] for r in data]).ravel()
        pid = int(m.group(1))
        out[pid] = out.get(pid, 0) + s
    return out


def convert(units_xlsx, plants_xlsx, load_dir, out_dir, system: str = "西南",
            accident_rate: float = 0.04, hydro_in_curtail_base: bool = True,
            security: bool = False, coal_must_run: bool = False) -> dict:
    wb = openpyxl.load_workbook(units_xlsx, data_only=True)
    sysrow = next(r for r in _table(wb, "系统表") if r["名称"] == system)
    sid = sysrow["系统ID"]
    loadrow = next(r for r in _table(wb, "负荷表") if r["系统ID"] == sid)
    nodes = {r["节点ID"] for r in _table(wb, "节点表") if r["系统ID"] == sid}
    plants = {r["电站ID"]: r for r in _table(wb, "电站表")
              if r["节点ID"] in nodes and r["有效性"] == 1}

    curve_name = sysrow["日历负荷"]
    curve_file = next(Path(load_dir).glob(f"*{curve_name}*.xlsx"))
    load = read_daily_curve(curve_file) * float(loadrow["年Pmax"])
    T = load.size

    active = set()                            # 有有效机组的电站
    units, storages = [], defaultdict(lambda: {"power": 0.0, "energy": 0.0, "eta": None})
    for u in _table(wb, "机组表"):
        p = plants.get(u["电站ID"])
        n = int(u.get("台数") or 0)
        if p is None or u["有效性"] != 1 or n <= 0:
            continue
        typ, cap = p["类型"], float(u["单机容量"])
        active.add(u["电站ID"])
        if typ in COAL | GAS | OTHER | NUCLEAR:
            coal = {**{t: float(u.get("燃料单耗") or 0) for t in COAL},
                    **{t: GAS_PRIORITY for t in GAS}, **{t: OTHER_PRIORITY for t in OTHER | NUCLEAR}}[typ]
            # 燃机在 GTDS 特性表中为启停调峰（最小出力 0），机组表"技术出力 1"不是最小出力
            pmin = 0.0 if typ in GAS else cap * float(u["技术出力"])
            for k in range(n):
                units.append(dict(name=f"{p['名称']}#{k + 1}", capacity=cap,
                                  pmin=round(pmin, 3), coal=coal,
                                  must_run=typ in OTHER | NUCLEAR | (COAL if coal_must_run else set()),
                                  must_run_weeks=""))
        elif typ in PUMPED | BATTERY:
            key = "抽水蓄能" if typ in PUMPED else "新型储能"
            s = storages[key]
            s["power"] += cap * n
            eta = float(p["储能效率"])
            # 储能库容为满功率充电时长；GTDS 发电库容 = 库容 × η（如抽蓄 7 h → 5.25 h）。
            # 本算法的容量是储能侧电量，充电侧 × η_c = 库容 × √η
            s["energy"] += cap * n * float(u["储能库容"]) * eta ** 0.5
            s["eta"] = eta

    series = read_plant_sheets(plants_xlsx)
    ts = {k: np.zeros(T) for k in ("wind", "solar", "hydro")}
    missing = []
    for pid, p in plants.items():
        kind = ("wind" if p["类型"] in WIND else "solar" if p["类型"] in SOLAR
                else "hydro" if p["类型"] in HYDRO else None)
        if kind is None or pid not in active:
            continue
        if pid in series:
            ts[kind] += series[pid][:T]
        else:
            missing.append(p["名称"])

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    pd.DataFrame({"load": load, **ts}).round(3).to_csv(out / "timeseries.csv", index=False)
    pd.DataFrame(units).to_csv(out / "units.csv", index=False)
    pd.DataFrame([dict(name=k, power=v["power"], energy=v["energy"], eta_rt=v["eta"],
                       soc_min=0.0, soc_max=1.0, self_discharge_day=0.0)
                  for k, v in storages.items() if v["power"] > 0]).to_csv(out / "storages.csv", index=False)
    cfg = {"load_reserve_rate": float(loadrow["负荷备用"]), "accident_reserve_rate": accident_rate,
           "hydro_in_curtail_base": hydro_in_curtail_base,
           # 系统表"保安开机"：在运火电容量下限占火电总装机的比例
           "min_online_ratio": float(sysrow["保安开机"] or 0) if security else 0.0}
    (out / "config.json").write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"系统": system, "负荷曲线": curve_name, "年最大负荷MW": float(loadrow["年Pmax"]),
            "火电台数": len(units), "火电容量MW": sum(u["capacity"] for u in units),
            "储能": {k: dict(v) for k, v in storages.items()},
            "风电TWh": ts["wind"].sum() / 1e6, "光伏TWh": ts["solar"].sum() / 1e6,
            "水电TWh": ts["hydro"].sum() / 1e6, "负荷TWh": load.sum() / 1e6,
            "缺少出力曲线的电站": missing}
